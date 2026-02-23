import hashlib
import re
import uuid
from datetime import date
from calendar import monthrange

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session


MONTH_MAP = {
    "JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6,
    "JUL": 7, "AUG": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12,
}

OFF_CODES = {"OFF", "O", "WO", "HOL", "H", "AL", "SL", "VL", "L", "หยุด", "ลา", "X", "HD", "V"}
OFFSITE_CODES = {"OFFSITE", "FIELD", "REMOTE", "WFH", "OS"}


def _sha256_bytes(b: bytes) -> str:
    h = hashlib.sha256()
    h.update(b)
    return h.hexdigest()


def _detect_month_year(ws) -> tuple[int, int]:
    """
    Try detect month/year from:
    - cell H1 like "Febuary'2026"
    - sheet title like "FEB"
    """
    # try H1 (row1 col8)
    v = ws.cell(1, 8).value
    if isinstance(v, str):
        s = v.strip()
        # "Febuary'2026" / "February 2026" / "Feb-2026"
        m = re.search(r"([A-Za-z]{3,})\D*([12]\d{3})", s)
        if m:
            mon = m.group(1).upper()[:3]
            yr = int(m.group(2))
            if mon in MONTH_MAP:
                return MONTH_MAP[mon], yr

    # fallback from sheet name
    title = (ws.title or "").upper().strip()
    title3 = title[:3]
    if title3 in MONTH_MAP:
        # year not in title -> read from H1 year if possible
        if isinstance(v, str):
            m2 = re.search(r"([12]\d{3})", v)
            if m2:
                return MONTH_MAP[title3], int(m2.group(1))
        raise ValueError("Cannot detect year from sheet. Put year in H1 (e.g. Febuary'2026).")

    raise ValueError("Cannot detect month/year from sheet. Expected sheet name like FEB and H1 contains year.")


def import_schedule_xlsx(db: Session, file_bytes: bytes, file_name: str, uploaded_by: str = "admin") -> dict:
    """
    Import schedule from the provided Excel (same structure as 02_Feb_2026.xlsx)
    Creates:
      - schedule_versions (1 row)
      - schedule_assignments (rows)
    Also auto-creates employees if emp_code not found (name from sheet).
    """
    sha = _sha256_bytes(file_bytes)

    # load workbook from bytes
    import io
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)

    # choose sheet: prefer a month sheet like FEB/JAN/... else first
    ws = None
    for name in wb.sheetnames:
        if name.upper()[:3] in MONTH_MAP:
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]

    month, year = _detect_month_year(ws)
    month_start = date(year, month, 1)
    last_day = monthrange(year, month)[1]

    # dynamic shift defaults from DB (allow changing default start times without code changes)
    shift_defs = db.execute(
        text("SELECT code, start_time FROM shift_definitions WHERE is_active=true")
    ).mappings().all()
    shift_start_map = {str(r["code"]).strip().upper(): r["start_time"] for r in shift_defs}
    work_shift_codes = set(shift_start_map.keys())

    # day columns are in row 3: numbers 1..31 at col 3..33
    day_cols = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(3, c).value
        if isinstance(v, int) and 1 <= v <= 31:
            day_cols.append((v, c))

    if not day_cols:
        raise ValueError("Cannot find day columns in row 3 (expected 1..31).")

    # prevent duplicate import of same file for same month (optional)
    dup = db.execute(
        text("SELECT 1 FROM schedule_versions WHERE month_start=:m AND file_sha256=:h LIMIT 1"),
        {"m": month_start, "h": sha},
    ).first()
    if dup:
        return {"ok": True, "message": "Already imported (same file hash)", "month_start": str(month_start)}

    ver_id = str(uuid.uuid4())
    db.execute(
        text("""
            INSERT INTO schedule_versions(id, month_start, source, uploaded_by, file_name, file_sha256)
            VALUES (:id, :m, 'excel', :u, :fn, :h)
        """),
        {"id": ver_id, "m": month_start, "u": uploaded_by, "fn": file_name, "h": sha},
    )

    inserted = 0
    created_employees = 0

    # employees start at row 4
    for r in range(4, ws.max_row + 1):
        emp_code = ws.cell(r, 1).value
        full_name = ws.cell(r, 2).value

        if emp_code is None and full_name is None:
            continue

        if emp_code is None:
            # skip malformed row
            continue

        emp_code_str = str(emp_code).strip()
        name_str = (str(full_name).strip() if full_name else emp_code_str)

        # ensure employee exists
        emp = db.execute(
            text("SELECT id FROM employees WHERE emp_code=:c LIMIT 1"),
            {"c": emp_code_str},
        ).mappings().first()

        if not emp:
            emp_id = str(uuid.uuid4())
            db.execute(
                text("INSERT INTO employees(id, emp_code, full_name, department, is_active) VALUES (:id,:c,:n,'',true)"),
                {"id": emp_id, "c": emp_code_str, "n": name_str},
            )
            created_employees += 1
        else:
            emp_id = emp["id"]

        # insert day by day
        for day, c in day_cols:
            if day > last_day:
                continue

            raw = ws.cell(r, c).value
            if raw is None or str(raw).strip() == "":
                continue

            code = str(raw).strip().upper()
            # normalize leave variants like HD* / V*
            normalized_code = code
            if normalized_code.startswith("HD"):
                normalized_code = "HD"
            elif normalized_code.startswith("V"):
                normalized_code = "V"

            work_date = date(year, month, day)

            is_workday = normalized_code in work_shift_codes
            is_off_code = normalized_code in OFF_CODES
            is_offsite_code = normalized_code in OFFSITE_CODES

            is_day_off = is_off_code
            work_mode = "OFFSITE" if is_offsite_code else "ONSITE"
            exempt = is_off_code or is_offsite_code
            note = ""

            start_time_override = shift_start_map.get(normalized_code) if is_workday else None

            if (not is_workday) and (not is_off_code) and (not is_offsite_code):
                # unknown code -> keep as exempt and note for manual review
                exempt = True
                note = f"UNMAPPED:{code}"
            elif is_off_code or is_offsite_code:
                note = code

            db.execute(
                text("""
                    INSERT INTO schedule_assignments(
                        id, version_id, employee_id, work_date, shift_code,
                        start_time_override, is_day_off, work_mode, exempt_attendance, note
                    )
                    VALUES(:id,:vid,:eid,:d,:sc,:sto,:off,:wm,:ex,:note)
                    ON CONFLICT (version_id, employee_id, work_date) DO UPDATE
                      SET shift_code=EXCLUDED.shift_code,
                          start_time_override=EXCLUDED.start_time_override,
                          is_day_off=EXCLUDED.is_day_off,
                          work_mode=EXCLUDED.work_mode,
                          exempt_attendance=EXCLUDED.exempt_attendance,
                          note=EXCLUDED.note
                """),
                {
                    "id": str(uuid.uuid4()),
                    "vid": ver_id,
                    "eid": emp_id,
                    "d": work_date,
                    "sc": normalized_code,
                    "sto": start_time_override,
                    "off": is_day_off,
                    "wm": work_mode,
                    "ex": exempt,
                    "note": note,
                },
            )
            inserted += 1

    db.commit()
    return {
        "ok": True,
        "month_start": str(month_start),
        "version_id": ver_id,
        "rows_inserted": inserted,
        "employees_created": created_employees,
        "file_sha256": sha,
        "sheet": ws.title,
    }
