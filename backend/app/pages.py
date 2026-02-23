import io
import os
import uuid
from datetime import date, datetime
from calendar import monthrange

try:
    import holidays as pyholidays
except Exception:
    pyholidays = None
from zoneinfo import ZoneInfo
from pathlib import Path
from fastapi import APIRouter, Request, Form, UploadFile, File
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text
from sqlalchemy.orm import Session
from fastapi import Depends
from .schedule_import import import_schedule_xlsx

from .db import get_db
from .config import settings
from .auth import require_login, is_logged_in

templates = Jinja2Templates(directory=str(Path(__file__).parent / "templates"))
router = APIRouter()

@router.get("/", response_class=HTMLResponse)
def home(request: Request):
    if is_logged_in(request):
        return RedirectResponse("/employees", status_code=302)
    return RedirectResponse("/login", status_code=302)

@router.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse("login.html", {"request": request, "error": None})

@router.post("/login")
def login(request: Request, username: str = Form(...), password: str = Form(...)):
    if username == settings.ADMIN_USER and password == settings.ADMIN_PASS:
        request.session["user"] = username
        return RedirectResponse("/employees", status_code=302)
    return templates.TemplateResponse("login.html", {"request": request, "error": "Invalid credentials"})

@router.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=302)

@router.get("/employees", response_class=HTMLResponse)
def employees_page(request: Request, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir: return redir

    rows = db.execute(text("SELECT id, emp_code, full_name, department, is_active, deleted_at, deleted_by FROM employees ORDER BY emp_code")).mappings().all()
    return templates.TemplateResponse("employees.html", {"request": request, "employees": rows, "admin_user": settings.ADMIN_USER})

@router.post("/employees")
def employees_add(
    request: Request,
    emp_code: str = Form(...),
    full_name: str = Form(...),
    department: str = Form(""),
    db: Session = Depends(get_db),
):
    redir = require_login(request)
    if redir: return redir

    db.execute(
        text("INSERT INTO employees(emp_code, full_name, department, is_active) VALUES (:c,:n,:d,true)"),
        {"c": emp_code.strip(), "n": full_name.strip(), "d": department.strip() or None},
    )
    db.commit()
    return RedirectResponse("/employees", status_code=302)


@router.post("/employees/{emp_id}/update")
def employees_update(
    request: Request,
    emp_id: str,
    emp_code: str = Form(...),
    full_name: str = Form(...),
    department: str = Form(""),
    is_active: str | None = Form(None),
    db: Session = Depends(get_db),
):
    redir = require_login(request)
    if redir:
        return redir

    db.execute(
        text(
            """
            UPDATE employees
            SET emp_code=:c,
                full_name=:n,
                department=:d,
                is_active=:a,
                deleted_at = CASE WHEN :a THEN NULL ELSE deleted_at END,
                deleted_by = CASE WHEN :a THEN NULL ELSE deleted_by END
            WHERE id=:id
            """
        ),
        {
            "id": emp_id,
            "c": emp_code.strip(),
            "n": full_name.strip(),
            "d": department.strip() or None,
            "a": bool(is_active),
        },
    )
    db.commit()
    return RedirectResponse("/employees", status_code=302)


@router.post("/employees/{emp_id}/delete")
def employees_delete(
    request: Request,
    emp_id: str,
    admin_password: str = Form(...),
    db: Session = Depends(get_db),
):
    redir = require_login(request)
    if redir:
        return redir

    if (admin_password or "") != settings.ADMIN_PASS:
        return RedirectResponse("/employees", status_code=302)

    db.execute(
        text(
            """
            UPDATE employees
            SET is_active=false, deleted_at=now(), deleted_by=:actor
            WHERE id=:id
            """
        ),
        {"id": emp_id, "actor": settings.ADMIN_USER},
    )
    db.execute(
        text("INSERT INTO audit_logs(actor, action, entity, entity_id, detail) VALUES (:a,:ac,:e,:id, CAST(:d AS jsonb))"),
        {"a": settings.ADMIN_USER, "ac": "SOFT_DELETE", "e": "employees", "id": emp_id, "d": '{"reason":"admin_confirm"}'},
    )
    db.commit()
    return RedirectResponse("/employees", status_code=302)

@router.get("/employees/{emp_id}/enroll", response_class=HTMLResponse)
def enroll_page(request: Request, emp_id: str, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir: return redir

    emp = db.execute(text("SELECT id, emp_code, full_name FROM employees WHERE id=:id"), {"id": emp_id}).mappings().first()
    if not emp:
        return RedirectResponse("/employees", status_code=302)

    session_id = str(uuid.uuid4())
    # store in session cookie (admin's browser)
    enroll_sessions = request.session.setdefault("enroll_sessions", {})
    enroll_sessions[session_id] = {"emp_id": emp_id, "count": 0}
    request.session["enroll_sessions"] = enroll_sessions

    return templates.TemplateResponse("enroll.html", {"request": request, "emp": emp, "session_id": session_id})

@router.post("/api/enroll/{session_id}/upload")
async def enroll_upload(request: Request, session_id: str, file: UploadFile = File(...)):
    redir = require_login(request)
    if redir: return redir

    enroll_sessions = request.session.get("enroll_sessions", {})
    s = enroll_sessions.get(session_id)
    if not s:
        return {"ok": False, "error": "invalid session"}

    # Save to disk
    tmp_dir = Path(__file__).resolve().parents[2] / "storage" / "enroll_tmp" / session_id
    tmp_dir.mkdir(parents=True, exist_ok=True)

    idx = int(s.get("count", 0)) + 1
    s["count"] = idx
    enroll_sessions[session_id] = s
    request.session["enroll_sessions"] = enroll_sessions

    out = tmp_dir / f"img_{idx:03d}.jpg"
    data = await file.read()
    out.write_bytes(data)

    return {"ok": True, "count": idx}

@router.post("/api/enroll/{session_id}/finalize")
def enroll_finalize(request: Request, session_id: str, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir: return redir

    enroll_sessions = request.session.get("enroll_sessions", {})
    s = enroll_sessions.get(session_id)
    if not s:
        return {"ok": False, "error": "invalid session"}

    emp_id = s["emp_id"]
    tmp_dir = Path(__file__).resolve().parents[2] / "storage" / "enroll_tmp" / session_id

    # Use backend-side embedder (optional) by importing from worker-style module if installed
    # We keep it lightweight: if insightface is available, we'll embed; otherwise error with instruction.
    try:
        from .enroll_embedder import embed_images_to_templates
    except Exception as e:
        return {"ok": False, "error": f"embedder import failed: {e}"}

    imgs = sorted(tmp_dir.glob("*.jpg"))
    if len(imgs) < 10:
        return {"ok": False, "error": "Please capture at least 10 images for enrollment."}

    embeddings, qualities = embed_images_to_templates([p.as_posix() for p in imgs])

    if len(embeddings) == 0:
        return {"ok": False, "error": "No valid faces detected in captured images."}

    # Keep top N templates by quality (max 5)
    pairs = sorted(list(zip(embeddings, qualities)), key=lambda x: x[1], reverse=True)[:5]

    # Determine embedding storage type (vector or float4[])
    has_vector = db.execute(text("SELECT EXISTS (SELECT 1 FROM pg_extension WHERE extname='vector') AS v")).mappings().first()["v"]

    if has_vector:
        # Insert as vector literal: '[...]'::vector
        for emb, q in pairs:
            db.execute(
                text("INSERT INTO face_templates(employee_id, embedding, quality_score) VALUES (:eid, (:emb)::vector, :q)"),
                {"eid": emp_id, "emb": "[" + ",".join(f"{x:.6f}" for x in emb) + "]", "q": float(q)},
            )
    else:
        # float4[]
        for emb, q in pairs:
            db.execute(
                text("INSERT INTO face_templates(employee_id, embedding, quality_score) VALUES (:eid, :emb, :q)"),
                {"eid": emp_id, "emb": list(map(float, emb)), "q": float(q)},
            )

    db.execute(
        text("INSERT INTO audit_logs(actor, action, entity, entity_id, detail) VALUES (:a, :ac, :e, :id, CAST(:d AS jsonb))"),
        {"a": settings.ADMIN_USER, "ac": "ENROLL_FACE", "e": "employees", "id": emp_id, "d": '{"source":"webcam"}'},
    )
    db.commit()

    # cleanup
    try:
        if tmp_dir.exists():
            for p in tmp_dir.glob("*"):
                p.unlink(missing_ok=True)
            tmp_dir.rmdir()
    except Exception:
        pass

    # remove session record
    enroll_sessions.pop(session_id, None)
    request.session["enroll_sessions"] = enroll_sessions

    return {"ok": True, "templates_saved": len(pairs)}

@router.get("/report/today", response_class=HTMLResponse)
def report_today(request: Request, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir: return redir

    rows = db.execute(text("""
        SELECT e.emp_code, e.full_name,
               (a.check_in_ts AT TIME ZONE 'Asia/Bangkok') AS check_in_th,
               COALESCE(a.attendance_status, 'ON_TIME') AS attendance_status,
               COALESCE(a.minutes_late, 0) AS minutes_late,
               a.policy_note
        FROM attendance_logs a
        JOIN employees e ON e.id = a.employee_id
        WHERE a.att_date = ((now() AT TIME ZONE 'Asia/Bangkok')::date)
        ORDER BY e.emp_code
    """)).mappings().all()

    return templates.TemplateResponse("report_today.html", {"request": request, "rows": rows, "tz": settings.TZ})

def _schedule_view_data(db: Session, month: str | None):
    if month:
        try:
            month_start = datetime.strptime(month, "%Y-%m").date().replace(day=1)
        except Exception:
            month_start = date.today().replace(day=1)
    else:
        v = db.execute(text("SELECT month_start FROM schedule_versions ORDER BY created_at DESC LIMIT 1")).scalar()
        month_start = v or date.today().replace(day=1)

    version = db.execute(
        text(
            """
            SELECT id, month_start, uploaded_by, file_name, created_at
            FROM schedule_versions
            WHERE month_start = :m
            ORDER BY created_at DESC
            LIMIT 1
            """
        ),
        {"m": month_start},
    ).mappings().first()

    rows = []
    if version:
        rows = db.execute(
            text(
                """
                SELECT sa.id, sa.work_date, sa.shift_code, sa.start_time_override,
                       sa.is_day_off, sa.work_mode, sa.exempt_attendance, sa.note,
                       e.emp_code, e.full_name
                FROM schedule_assignments sa
                JOIN employees e ON e.id = sa.employee_id
                WHERE sa.version_id = :vid
                ORDER BY sa.work_date, e.emp_code
                LIMIT 500
                """
            ),
            {"vid": version["id"]},
        ).mappings().all()

    shifts = db.execute(text("SELECT code, name FROM shift_definitions WHERE is_active=true ORDER BY code")).mappings().all()
    shift_codes = [s["code"] for s in shifts]
    return {
        "month": month_start.strftime("%Y-%m"),
        "version": version,
        "schedule_rows": rows,
        "shift_codes": shift_codes,
    }


@router.get("/schedule", response_class=HTMLResponse)
def schedule_page(request: Request, month: str | None = None, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir:
        return redir

    data = _schedule_view_data(db, month)
    return templates.TemplateResponse(
        "schedule_upload.html",
        {
            "request": request,
            "msg": None,
            "result": None,
            "default_path": settings.SCHEDULE_TEMPLATE_PATH,
            **data,
        },
    )


@router.get("/schedule/matrix", response_class=HTMLResponse)
def schedule_matrix(request: Request, month: str | None = None, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir:
        return redir

    if month:
        try:
            month_start = datetime.strptime(month, "%Y-%m").date().replace(day=1)
        except Exception:
            month_start = date.today().replace(day=1)
    else:
        v = db.execute(text("SELECT month_start FROM schedule_versions ORDER BY created_at DESC LIMIT 1")).scalar()
        month_start = v or date.today().replace(day=1)

    version = db.execute(
        text(
            """
            SELECT id, month_start, uploaded_by, file_name, created_at
            FROM schedule_versions
            WHERE month_start = :m
            ORDER BY created_at DESC
            LIMIT 1
            """
        ),
        {"m": month_start},
    ).mappings().first()

    day_count = monthrange(month_start.year, month_start.month)[1]

    thai_holiday_map: dict[int, str] = {}
    if pyholidays is not None:
        try:
            th_holidays = pyholidays.country_holidays("TH", years=[month_start.year], language="th")
            for d, name in th_holidays.items():
                if d.year == month_start.year and d.month == month_start.month:
                    thai_holiday_map[int(d.day)] = str(name)
        except Exception:
            thai_holiday_map = {}

    weekday_th = ["จ", "อ", "พ", "พฤ", "ศ", "ส", "อา"]
    day_meta = []
    for d in range(1, day_count + 1):
        dd = date(month_start.year, month_start.month, d)
        w = dd.weekday()  # Mon=0..Sun=6
        is_weekend = w >= 5
        holiday_name = thai_holiday_map.get(d)
        day_meta.append(
            {
                "day": d,
                "weekday": weekday_th[w],
                "is_weekend": is_weekend,
                "holiday_name": holiday_name,
                "is_holiday": holiday_name is not None,
            }
        )

    employees = db.execute(
        text(
            """
            SELECT id, emp_code, full_name
            FROM employees
            WHERE is_active=true
            ORDER BY emp_code
            """
        )
    ).mappings().all()

    matrix = []
    if version:
        rows = db.execute(
            text(
                """
                SELECT sa.employee_id, EXTRACT(DAY FROM sa.work_date)::int AS day_no,
                       sa.shift_code,
                       sa.note
                FROM schedule_assignments sa
                WHERE sa.version_id = :vid
                """
            ),
            {"vid": version["id"]},
        ).mappings().all()

        by_emp_day = {}
        for r in rows:
            key = (str(r["employee_id"]), int(r["day_no"]))
            code = (r.get("shift_code") or "").strip().upper()
            if not code:
                note = (r.get("note") or "").strip().upper()
                if note.startswith("UNMAPPED:"):
                    code = note.split(":", 1)[1].strip()
                else:
                    code = note
            by_emp_day[key] = code

        for e in employees:
            day_cells = []
            for dm in day_meta:
                d = dm["day"]
                code = by_emp_day.get((str(e["id"]), d), "")
                day_cells.append(
                    {
                        "day": d,
                        "code": code,
                        "is_weekend": dm["is_weekend"],
                        "is_holiday": dm["is_holiday"],
                        "holiday_name": dm["holiday_name"],
                    }
                )
            matrix.append({
                "emp_code": e["emp_code"],
                "full_name": e["full_name"],
                "cells": day_cells,
            })

    return templates.TemplateResponse(
        "schedule_matrix.html",
        {
            "request": request,
            "month": month_start.strftime("%Y-%m"),
            "day_meta": day_meta,
            "rows": matrix,
            "version": version,
        },
    )


@router.post("/schedule/upload", response_class=HTMLResponse)
async def schedule_upload(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir:
        return redir

    b = await file.read()
    try:
        result = import_schedule_xlsx(db, b, file.filename or "schedule.xlsx", uploaded_by=settings.ADMIN_USER)
        msg = "Import สำเร็จ"
        month = (result or {}).get("month_start", "")[:7] if result else None
    except Exception as e:
        result = None
        month = None
        msg = f"Import ไม่สำเร็จ: {e}"

    data = _schedule_view_data(db, month)
    return templates.TemplateResponse(
        "schedule_upload.html",
        {
            "request": request,
            "msg": msg,
            "result": result,
            "default_path": settings.SCHEDULE_TEMPLATE_PATH,
            **data,
        },
    )


@router.post("/schedule/import-default", response_class=HTMLResponse)
def schedule_import_default(request: Request, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir:
        return redir

    p = Path(settings.SCHEDULE_TEMPLATE_PATH)
    if not p.exists():
        data = _schedule_view_data(db, None)
        return templates.TemplateResponse(
            "schedule_upload.html",
            {
                "request": request,
                "msg": f"ไม่พบไฟล์ template: {p}",
                "result": None,
                "default_path": settings.SCHEDULE_TEMPLATE_PATH,
                **data,
            },
        )

    try:
        b = p.read_bytes()
        result = import_schedule_xlsx(db, b, p.name, uploaded_by=settings.ADMIN_USER)
        msg = f"Import template สำเร็จ: {p}"
        month = (result or {}).get("month_start", "")[:7] if result else None
    except Exception as e:
        result = None
        month = None
        msg = f"Import template ไม่สำเร็จ: {e}"

    data = _schedule_view_data(db, month)
    return templates.TemplateResponse(
        "schedule_upload.html",
        {
            "request": request,
            "msg": msg,
            "result": result,
            "default_path": settings.SCHEDULE_TEMPLATE_PATH,
            **data,
        },
    )


@router.post("/schedule/assignment/{assignment_id}/update", response_class=HTMLResponse)
def schedule_assignment_update(
    request: Request,
    assignment_id: str,
    month: str = Form(...),
    shift_code: str = Form(""),
    start_time_override: str = Form(""),
    work_mode: str = Form("ONSITE"),
    is_day_off: str | None = Form(None),
    exempt_attendance: str | None = Form(None),
    note: str = Form(""),
    db: Session = Depends(get_db),
):
    redir = require_login(request)
    if redir:
        return redir

    db.execute(
        text(
            """
            UPDATE schedule_assignments
            SET shift_code=:sc,
                start_time_override=:sto,
                work_mode=:wm,
                is_day_off=:off,
                exempt_attendance=:ex,
                note=:note
            WHERE id=:id
            """
        ),
        {
            "id": assignment_id,
            "sc": (shift_code or "").strip().upper() or None,
            "sto": (start_time_override or "").strip() or None,
            "wm": (work_mode or "ONSITE").strip().upper(),
            "off": bool(is_day_off),
            "ex": bool(exempt_attendance),
            "note": (note or "").strip() or None,
        },
    )
    db.commit()

    data = _schedule_view_data(db, month)
    return templates.TemplateResponse(
        "schedule_upload.html",
        {
            "request": request,
            "msg": "บันทึกการแก้ไข schedule แล้ว",
            "result": None,
            "default_path": settings.SCHEDULE_TEMPLATE_PATH,
            **data,
        },
    )

@router.get("/shifts", response_class=HTMLResponse)
def shifts_page(request: Request, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir:
        return redir

    rows = db.execute(
        text(
            """
            SELECT code, name, start_time, grace_minutes, absent_after_minutes, is_active
            FROM shift_definitions
            ORDER BY code
            """
        )
    ).mappings().all()

    return templates.TemplateResponse(
        "shift_settings.html",
        {"request": request, "rows": rows, "msg": None},
    )


@router.post("/shifts/{code}/update", response_class=HTMLResponse)
def shifts_update(
    request: Request,
    code: str,
    start_time: str = Form(...),
    grace_minutes: int = Form(15),
    absent_after_minutes: int = Form(120),
    is_active: str | None = Form(None),
    db: Session = Depends(get_db),
):
    redir = require_login(request)
    if redir:
        return redir

    code_norm = (code or "").strip().upper()
    db.execute(
        text(
            """
            UPDATE shift_definitions
            SET start_time = :st,
                grace_minutes = :gm,
                absent_after_minutes = :am,
                is_active = :ia,
                updated_at = now()
            WHERE code = :code
            """
        ),
        {
            "code": code_norm,
            "st": (start_time or "").strip(),
            "gm": int(grace_minutes),
            "am": int(absent_after_minutes),
            "ia": bool(is_active),
        },
    )
    db.commit()

    rows = db.execute(
        text(
            """
            SELECT code, name, start_time, grace_minutes, absent_after_minutes, is_active
            FROM shift_definitions
            ORDER BY code
            """
        )
    ).mappings().all()

    return templates.TemplateResponse(
        "shift_settings.html",
        {"request": request, "rows": rows, "msg": f"บันทึกค่า Shift {code_norm} แล้ว"},
    )


@router.get("/monitor/cameras", response_class=HTMLResponse)
def monitor_cameras(request: Request, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir:
        return redir

    rows = db.execute(
        text(
            """
            SELECT
              c.name AS camera_name,
              c.location,
              c.is_enabled,
              ws.worker_name,
              ws.status,
              ws.last_seen,
              ws.last_error,
              ws.preview_path,
              (
                SELECT fe.event_ts
                FROM face_events fe
                WHERE fe.camera_id = c.id
                ORDER BY fe.event_ts DESC
                LIMIT 1
              ) AS last_event_ts
            FROM cameras c
            LEFT JOIN camera_worker_status ws ON ws.camera_id = c.id
            ORDER BY c.name
            """
        )
    ).mappings().all()

    now_local = datetime.now(ZoneInfo(settings.TZ))

    def _to_url(p: str | None) -> str | None:
        if not p:
            return None
        s = str(p).replace('\\', '/')
        mark = '/storage/'
        i = s.lower().find(mark)
        if i >= 0:
            return s[i:]
        i2 = s.lower().find('storage/')
        if i2 >= 0:
            return '/' + s[i2:]
        return None

    vm = []
    for r in rows:
        d = dict(r)
        d["preview_url"] = _to_url(d.get("preview_path"))
        vm.append(d)

    return templates.TemplateResponse(
        "camera_monitor.html",
        {"request": request, "rows": vm, "now_local": now_local, "tz": settings.TZ},
    )


@router.get("/report/monthly", response_class=HTMLResponse)
def report_monthly(request: Request, month: str | None = None, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir:
        return redir

    if month:
        try:
            month_start = datetime.strptime(month, "%Y-%m").date().replace(day=1)
        except Exception:
            month_start = date.today().replace(day=1)
    else:
        month_start = date.today().replace(day=1)

    rows = db.execute(
        text(
            """
            SELECT e.emp_code, e.full_name,
                   COALESCE(SUM(CASE WHEN a.attendance_status='ON_TIME' THEN 1 ELSE 0 END),0) AS on_time_count,
                   COALESCE(SUM(CASE WHEN a.attendance_status='LATE' THEN 1 ELSE 0 END),0) AS late_count,
                   COALESCE(SUM(CASE WHEN a.attendance_status='ABSENT' THEN 1 ELSE 0 END),0) AS absent_count,
                   COALESCE(SUM(CASE WHEN a.attendance_status='EXEMPT' THEN 1 ELSE 0 END),0) AS exempt_count,
                   COALESCE(SUM(CASE WHEN a.attendance_status='OFF' THEN 1 ELSE 0 END),0) AS off_count,
                   COALESCE(SUM(CASE WHEN COALESCE(a.minutes_late,0) > 0 THEN a.minutes_late ELSE 0 END),0) AS late_minutes_total
            FROM employees e
            LEFT JOIN attendance_logs a
              ON a.employee_id = e.id
             AND date_trunc('month', a.att_date)::date = :m
            WHERE e.is_active = true
            GROUP BY e.emp_code, e.full_name
            ORDER BY e.emp_code
            """
        ),
        {"m": month_start},
    ).mappings().all()

    top_late = db.execute(
        text(
            """
            SELECT e.emp_code, e.full_name,
                   COALESCE(SUM(CASE WHEN a.attendance_status='LATE' THEN 1 ELSE 0 END),0) AS late_count
            FROM employees e
            LEFT JOIN attendance_logs a
              ON a.employee_id=e.id
             AND date_trunc('month', a.att_date)::date=:m
            WHERE e.is_active=true
            GROUP BY e.emp_code, e.full_name
            ORDER BY late_count DESC, e.emp_code
            LIMIT 5
            """
        ),
        {"m": month_start},
    ).mappings().all()

    top_absent = db.execute(
        text(
            """
            SELECT e.emp_code, e.full_name,
                   COALESCE(SUM(CASE WHEN a.attendance_status='ABSENT' THEN 1 ELSE 0 END),0) AS absent_count
            FROM employees e
            LEFT JOIN attendance_logs a
              ON a.employee_id=e.id
             AND date_trunc('month', a.att_date)::date=:m
            WHERE e.is_active=true
            GROUP BY e.emp_code, e.full_name
            ORDER BY absent_count DESC, e.emp_code
            LIMIT 5
            """
        ),
        {"m": month_start},
    ).mappings().all()

    return templates.TemplateResponse(
        "report_monthly.html",
        {
            "request": request,
            "rows": rows,
            "month": month_start.strftime("%Y-%m"),
            "top_late": top_late,
            "top_absent": top_absent,
        },
    )


@router.get("/report/monthly/export")
def report_monthly_export(request: Request, month: str | None = None, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir:
        return redir

    if month:
        try:
            month_start = datetime.strptime(month, "%Y-%m").date().replace(day=1)
        except Exception:
            month_start = date.today().replace(day=1)
    else:
        month_start = date.today().replace(day=1)

    rows = db.execute(
        text(
            """
            SELECT e.emp_code, e.full_name,
                   COALESCE(SUM(CASE WHEN a.attendance_status='ON_TIME' THEN 1 ELSE 0 END),0) AS on_time_count,
                   COALESCE(SUM(CASE WHEN a.attendance_status='LATE' THEN 1 ELSE 0 END),0) AS late_count,
                   COALESCE(SUM(CASE WHEN a.attendance_status='ABSENT' THEN 1 ELSE 0 END),0) AS absent_count,
                   COALESCE(SUM(CASE WHEN a.attendance_status='EXEMPT' THEN 1 ELSE 0 END),0) AS exempt_count,
                   COALESCE(SUM(CASE WHEN a.attendance_status='OFF' THEN 1 ELSE 0 END),0) AS off_count,
                   COALESCE(SUM(CASE WHEN COALESCE(a.minutes_late,0) > 0 THEN a.minutes_late ELSE 0 END),0) AS late_minutes_total
            FROM employees e
            LEFT JOIN attendance_logs a
              ON a.employee_id = e.id
             AND date_trunc('month', a.att_date)::date = :m
            WHERE e.is_active = true
            GROUP BY e.emp_code, e.full_name
            ORDER BY e.emp_code
            """
        ),
        {"m": month_start},
    ).mappings().all()

    # detail sheet rows (daily)
    detail_rows = db.execute(
        text(
            """
            SELECT e.emp_code, e.full_name, a.att_date,
                   (a.check_in_ts AT TIME ZONE 'Asia/Bangkok') AS check_in_th,
                   COALESCE(a.shift_code,'') AS shift_code,
                   COALESCE(a.attendance_status,'ON_TIME') AS attendance_status,
                   COALESCE(a.minutes_late,0) AS minutes_late,
                   COALESCE(a.policy_note,'') AS policy_note
            FROM attendance_logs a
            JOIN employees e ON e.id = a.employee_id
            WHERE date_trunc('month', a.att_date)::date = :m
            ORDER BY a.att_date, e.emp_code
            """
        ),
        {"m": month_start},
    ).mappings().all()

    wb = Workbook()

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Summary"

    title = f"Attendance Summary {month_start.strftime('%Y-%m')}"
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["Emp Code", "Name", "ON_TIME", "LATE", "ABSENT", "EXEMPT", "OFF", "Late Minutes"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for c in range(1, 9):
        cell = ws.cell(row=2, column=c)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([
            r["emp_code"],
            r["full_name"],
            int(r["on_time_count"]),
            int(r["late_count"]),
            int(r["absent_count"]),
            int(r["exempt_count"]),
            int(r["off_count"]),
            int(r["late_minutes_total"]),
        ])

    for col, width in {"A": 14, "B": 28, "C": 10, "D": 10, "E": 10, "F": 10, "G": 10, "H": 14}.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=3, min_col=3, max_col=8):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")

    # Sheet 2: Daily Detail
    wd = wb.create_sheet("Detail")
    wd_headers = ["Date", "Emp Code", "Name", "Shift", "Check-in (TH)", "Status", "Late(min)", "Note"]
    wd.append(wd_headers)

    for c in range(1, 9):
        cell = wd.cell(row=1, column=c)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for r in detail_rows:
        wd.append([
            r["att_date"],
            r["emp_code"],
            r["full_name"],
            r["shift_code"],
            r["check_in_th"],
            r["attendance_status"],
            int(r["minutes_late"] or 0),
            r["policy_note"],
        ])

    for col, width in {"A": 12, "B": 14, "C": 28, "D": 8, "E": 22, "F": 12, "G": 10, "H": 36}.items():
        wd.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    fname = f"attendance_summary_{month_start.strftime('%Y_%m')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )

@router.get("/dashboard/monthly", response_class=HTMLResponse)
def dashboard_monthly(request: Request, month: str | None = None, db: Session = Depends(get_db)):
    redir = require_login(request)
    if redir:
        return redir

    if month:
        try:
            month_start = datetime.strptime(month, "%Y-%m").date().replace(day=1)
        except Exception:
            month_start = date.today().replace(day=1)
    else:
        month_start = date.today().replace(day=1)

    now_local = datetime.now(ZoneInfo(settings.TZ))
    today_local = now_local.date()

    kpi = db.execute(text("""
        SELECT
          COALESCE(SUM(CASE WHEN attendance_status='LATE' THEN 1 ELSE 0 END),0) AS total_late,
          COALESCE(SUM(CASE WHEN attendance_status='ABSENT' THEN 1 ELSE 0 END),0) AS total_absent,
          COALESCE(SUM(CASE WHEN attendance_status='EXEMPT' THEN 1 ELSE 0 END),0) AS total_exempt,
          COALESCE(SUM(CASE WHEN attendance_status='ON_TIME' THEN 1 ELSE 0 END),0) AS total_on_time,
          COALESCE(SUM(CASE WHEN attendance_status='OFF' THEN 1 ELSE 0 END),0) AS total_off
        FROM attendance_logs
        WHERE date_trunc('month', att_date)::date = :m
    """), {"m": month_start}).mappings().first()

    top_late = db.execute(text("""
        SELECT e.emp_code, e.full_name, COUNT(*)::int AS late_count
        FROM attendance_logs a
        JOIN employees e ON e.id=a.employee_id
        WHERE date_trunc('month', a.att_date)::date=:m AND a.attendance_status='LATE'
        GROUP BY e.emp_code, e.full_name
        ORDER BY late_count DESC, e.emp_code
        LIMIT 10
    """), {"m": month_start}).mappings().all()

    top_absent = db.execute(text("""
        SELECT e.emp_code, e.full_name, COUNT(*)::int AS absent_count
        FROM attendance_logs a
        JOIN employees e ON e.id=a.employee_id
        WHERE date_trunc('month', a.att_date)::date=:m AND a.attendance_status='ABSENT'
        GROUP BY e.emp_code, e.full_name
        ORDER BY absent_count DESC, e.emp_code
        LIMIT 10
    """), {"m": month_start}).mappings().all()

    live_rows = db.execute(
        text(
            """
            WITH latest_version AS (
              SELECT id
              FROM schedule_versions
              WHERE month_start = date_trunc('month', CAST(:today AS date))::date
              ORDER BY created_at DESC
              LIMIT 1
            )
            SELECT
              COALESCE(sa.shift_code, '-') AS shift_code,
              e.emp_code,
              e.full_name,
              COALESCE(sa.start_time_override, sd.start_time) AS shift_start,
              (
                SELECT fe.snapshot_path
                FROM face_events fe
                WHERE fe.match_employee_id = e.id
                  AND (fe.event_ts AT TIME ZONE 'Asia/Bangkok')::date = :today
                  AND fe.snapshot_path IS NOT NULL
                ORDER BY fe.event_ts DESC
                LIMIT 1
              ) AS snapshot_path,
              COALESCE(a.attendance_status,
                CASE
                  WHEN sa.is_day_off THEN 'OFF'
                  WHEN sa.exempt_attendance OR UPPER(COALESCE(sa.work_mode,'ONSITE')) IN ('OFFSITE','REMOTE','FIELD') THEN 'EXEMPT'
                  ELSE 'PENDING'
                END
              ) AS live_status,
              COALESCE(a.minutes_late, 0) AS minutes_late,
              (a.check_in_ts AT TIME ZONE 'Asia/Bangkok') AS check_in_local
            FROM schedule_assignments sa
            JOIN latest_version lv ON lv.id = sa.version_id
            JOIN employees e ON e.id = sa.employee_id AND e.is_active = true
            LEFT JOIN shift_definitions sd ON sd.code = sa.shift_code
            LEFT JOIN attendance_logs a ON a.employee_id = sa.employee_id AND a.att_date = sa.work_date
            WHERE sa.work_date = :today
              AND COALESCE(sa.start_time_override, sd.start_time) IS NOT NULL
              AND :now_local >= (CAST(:today AS timestamp) + COALESCE(sa.start_time_override, sd.start_time))
              AND :now_local < ((CAST(:today AS timestamp) + COALESCE(sa.start_time_override, sd.start_time)) + interval '9 hour')
            ORDER BY sa.shift_code, e.emp_code
            """
        ),
        {"today": today_local, "now_local": now_local.replace(tzinfo=None)},
    ).mappings().all()

    def _to_snapshot_url(p: str | None) -> str | None:
        if not p:
            return None
        s = str(p).replace('\\', '/')
        mark = '/storage/'
        i = s.lower().find(mark)
        if i >= 0:
            return s[i:]
        i2 = s.lower().find('storage/')
        if i2 >= 0:
            return '/' + s[i2:]
        return None

    live_by_shift = {}
    live_summary = {"ON_TIME": 0, "LATE": 0, "ABSENT": 0, "EXEMPT": 0, "OFF": 0, "PENDING": 0}
    for r in live_rows:
        d = dict(r)
        d["snapshot_url"] = _to_snapshot_url(d.get("snapshot_path"))
        sc = d.get("shift_code") or "-"
        live_by_shift.setdefault(sc, []).append(d)
        st = (d.get("live_status") or "PENDING").upper()
        if st not in live_summary:
            live_summary[st] = 0
        live_summary[st] += 1

    return templates.TemplateResponse("dashboard_monthly.html", {
        "request": request,
        "month": month_start.strftime("%Y-%m"),
        "kpi": kpi,
        "top_late": top_late,
        "top_absent": top_absent,
        "live_by_shift": live_by_shift,
        "live_summary": live_summary,
        "now_local": now_local,
    })
